"""Read-only mechanical checks for the four personal development skills.

Canonical repository source. Each skill ships an identical portable copy.
No business decisions, automatic fixes, writes, builds, or semantic API claims.
"""
import argparse
import datetime as dt
import json
import re
import sys
from collections import deque
from pathlib import Path
from urllib.parse import unquote

SCHEMAS = {
    'documentation': ['system', 'side', 'updated', 'code_revision', 'sources', 'dependencies'],
    'planning': ['id', 'kind', 'system', 'title', 'status', 'updated', 'sources', 'links', 'code_revision'],
    'skill': ['name', 'description'],
}
SIDES = {'client', 'server', 'shared'}


def output(value):
    print(json.dumps(value, ensure_ascii=False, indent=2, default=str))


def yaml_meta(text):
    import yaml

    class UniqueLoader(yaml.SafeLoader):
        pass

    def mapping(loader, node, deep=False):
        result = {}
        for key_node, value_node in node.value:
            key = loader.construct_object(key_node, deep=deep)
            if key in result:
                raise ValueError('duplicate YAML key: ' + str(key))
            result[key] = loader.construct_object(value_node, deep=deep)
        return result

    UniqueLoader.add_constructor(yaml.resolver.BaseResolver.DEFAULT_MAPPING_TAG, mapping)
    match = re.match(r'\A---\r?\n(.*?)\r?\n---(?:\r?\n|$)', text, re.S)
    if not match:
        raise ValueError('missing YAML frontmatter')
    data = yaml.load(match[1], Loader=UniqueLoader)
    if not isinstance(data, dict):
        raise ValueError('meta must be a mapping')
    return data, text[match.end():], match.end()


def slug(title):
    title = re.sub(r'[^\w\s-]', '', title.lower(), flags=re.UNICODE)
    return title.replace(' ', '-')


def check(args):
    path = Path(args.path).resolve()
    if not path.exists():
        raise ValueError('target does not exist')
    candidates = [path] if path.is_file() else sorted(path.rglob('*'))
    issues, count, skipped = [], 0, []

    def fail(file, code, message, line=1):
        issues.append({'file': str(file), 'line': line, 'rule': code, 'message': message})

    def link(file, ref, line=1):
        if ref.startswith(('https://', 'http://', 'mailto:')):
            return  # Remote existence is deliberately not checked.
        ref = unquote(ref.strip('<>'))
        target_name, _, anchor = ref.partition('#')
        target = file.parent / target_name if target_name else file
        if not target.exists():
            fail(file, 'LINK', 'missing target: ' + ref, line)
            return
        if anchor and target.suffix.lower() == '.md':
            headers = re.findall(r'^#{1,6}\s+(.+)$', target.read_text(encoding='utf-8-sig'), re.M)
            if anchor not in {slug(h) for h in headers}:
                fail(file, 'ANCHOR', 'missing Markdown heading: ' + ref, line)

    for file in candidates:
        if not file.is_file() or '.git' in file.parts:
            continue
        if 'templates' in file.parts:
            skipped.append(str(file))
            continue
        in_api = 'api' in file.parts and 'references' not in file.parts
        if in_api and file.suffix.lower() in {'.md', '.json', '.yaml', '.yml'}:
            fail(file, 'API_DIRECTORY', 'API directory contains a non-declaration file')
            continue
        if in_api and re.search(r'\.api\.(cs|d\.ts|pyi)$', file.name):
            count += 1
            text = file.read_text(encoding='utf-8-sig')
            relative = file.parts[file.parts.index('api') + 1:]
            if len(relative) < 3 or relative[0] not in SIDES:
                fail(file, 'API_SIDE', 'expected api/client|server|shared/module/file')
            for field in ['System', 'Side', 'Language', 'Source', 'CodeRevision', 'Status', 'Updated', 'Verification']:
                if not re.search(r'(?m)^\s*(?://|#)\s*' + field + r':\s*\S', text):
                    fail(file, 'API_HEADER', 'missing header: ' + field)
            side = re.search(r'(?m)^\s*(?://|#)\s*Side:\s*(\w+)', text)
            if side and relative and side[1] != relative[0]:
                fail(file, 'API_SIDE', 'header Side differs from directory')
            status = re.search(r'(?m)^\s*(?://|#)\s*Status:\s*(\w+)', text)
            if status and status[1] not in {'planned', 'mixed', 'implemented'}:
                fail(file, 'API_STATUS', 'invalid declaration status')
            if status and '文档' in file.parts and status[1] != 'implemented':
                fail(file, 'API_HANDOFF', 'unfinished declaration placed in formal docs')
            # Only simple C# structural mistakes; this is not Roslyn/TypeScript/Python API extraction.
            if file.suffix == '.cs':
                code = re.sub(r'(?m)^\s*//.*$', '', text)
                if re.search(r'\bpublic\s+(?:static\s+|readonly\s+)*[\w<>\[\]?]+\s+\w+\s*(?:;|=|\{\s*(?:get|set))', code):
                    fail(file, 'API_MEMBER', 'field/property found in public-function declaration')
                if re.search(r'\bpublic\s+[^;{}\n]+\([^;{}]*\)\s*\{', code):
                    fail(file, 'API_BODY', 'implementation body found')
            continue
        if file.suffix.lower() != '.md':
            continue
        count += 1
        text = file.read_text(encoding='utf-8-sig')
        try:
            meta, body, offset = yaml_meta(text)
        except Exception as exc:
            fail(file, 'META', str(exc))
            continue
        kind = args.kind
        if file.name == 'SKILL.md' or ('name' in meta and 'description' in meta):
            kind = 'skill'
        elif kind == 'auto':
            kind = 'documentation' if 'side' in meta else 'planning'
        for field in SCHEMAS[kind]:
            if field not in meta:
                fail(file, 'REQUIRED', 'missing meta field: ' + field)
        string_fields = {'name', 'description'} if kind == 'skill' else {'system', 'code_revision'}
        if kind == 'planning':
            string_fields |= {'id', 'title', 'kind', 'status'}
        for field in string_fields & meta.keys():
            if not isinstance(meta[field], str) or not meta[field].strip():
                fail(file, 'TYPE', field + ' must be a nonempty string')
        if kind != 'skill':
            try:
                date = str(meta.get('updated', ''))
                if not re.fullmatch(r'\d{4}-\d{2}-\d{2}', date):
                    raise ValueError()
                dt.date.fromisoformat(date)
            except ValueError:
                fail(file, 'DATE', 'updated must be a real YYYY-MM-DD date')
            sources = meta.get('sources')
            if not isinstance(sources, list) or any(not isinstance(s, str) or not s for s in sources):
                fail(file, 'SOURCES', 'sources must be a list of nonempty relative paths')
            else:
                for source in sources:
                    if re.match(r'^(?:[A-Za-z]:|/|\\)', source):
                        fail(file, 'SOURCE_PATH', 'source path must be relative: ' + source)
                    elif args.project_root:
                        base = Path(args.project_root).resolve()
                        resolved = (base / source).resolve()
                        if not resolved.is_relative_to(base) or not resolved.exists():
                            fail(file, 'SOURCE_PATH', 'missing/outside project source: ' + source)
        if kind == 'planning':
            if meta.get('kind') not in {'index', 'plan', 'system'}:
                fail(file, 'KIND', 'invalid planning kind')
            if meta.get('status') not in {'draft', 'ready', 'implemented', 'superseded', 'cancelled', 'partial'}:
                fail(file, 'STATUS', 'invalid planning status')
            refs = meta.get('links', [])
            if not isinstance(refs, list):
                fail(file, 'LINKS', 'links must be a list')
            else:
                for ref in refs:
                    if isinstance(ref, str):
                        link(file, ref)
                    else:
                        fail(file, 'LINKS', 'link must be a string')
            if file.name == 'config.md':
                section = re.search(r'(?m)^##\s+(?:Excel\s*)?字段字典\s*\n(.*?)(?=^##\s|\Z)', body, re.S | re.M)
                if not section:
                    fail(file, 'CONFIG_FIELDS', 'missing field dictionary section')
                else:
                    table = [line.strip() for line in section[1].splitlines() if line.strip().startswith('|')]
                    if len(table) < 3:
                        fail(file, 'CONFIG_FIELDS', 'field dictionary has no data rows')
                    else:
                        cells = lambda row: [c.strip() for c in re.split(r'(?<!\\)\|', row.strip('|'))]
                        headers = cells(table[0])
                        comments = [i for i, h in enumerate(headers) if '注释' in h or 'TIPS' in h]
                        if not comments:
                            fail(file, 'CONFIG_TIPS', 'dictionary must include a field-comment/TIPS column')
                        for row in table[2:]:
                            values = cells(row)
                            if len(values) != len(headers):
                                fail(file, 'CONFIG_COLUMNS', 'field row/header column counts differ')
                            elif comments and not values[comments[0]]:
                                fail(file, 'CONFIG_TIPS', 'field has an empty comment: ' + values[0])
                if not re.search(r'(?m)^##\s+.*处理清单', body):
                    fail(file, 'CONFIG_ACTIONS', 'missing existing/missing-field processing list')
        if kind == 'documentation':
            extra = set(meta) - set(SCHEMAS[kind]) - {'version'} - set(args.allow_meta_field)
            if extra:
                fail(file, 'COMPACT_META', 'unexpected fields: ' + ', '.join(sorted(extra)))
            sides = meta.get('side')
            if not isinstance(sides, list) or not sides or any(s not in SIDES for s in sides):
                fail(file, 'SIDE', 'side must be a nonempty client/server/shared list')
            deps = meta.get('dependencies')
            if not isinstance(deps, list):
                fail(file, 'DEPENDENCIES', 'dependencies must be a list')
            else:
                for dep in deps:
                    if not isinstance(dep, dict) or not isinstance(dep.get('document'), str) or not isinstance(dep.get('uses'), str) or not dep.get('uses', '').strip():
                        fail(file, 'DEPENDENCY', 'expected document and nonempty uses')
                        continue
                    if Path(dep['document'].split('#')[0]).suffix != '.md':
                        fail(file, 'DEPENDENCY', 'dependency must map to a Markdown system document')
                    if re.search(r'(?i)(?:^|/)(?:verification|test-report|report|tests?)(?:/|\.md)', dep['document']):
                        fail(file, 'TEST_DEPENDENCY', 'test report is not a system dependency')
                    link(file, dep['document'])
            headings = re.findall(r'^#{1,6}\s+(.+)$', body, re.M)
            for heading in headings:
                if re.search(r'测试|验收|验证结果|核对与缺口|执行记录|(?i:testing|test results)', heading):
                    fail(file, 'TEST_SECTION', 'test content must be in a separate report: ' + heading)
            graphs = re.findall(r'```mermaid\s*\n(.*?)\n```', body, re.S)
            if not args.topic and not graphs and not re.search(r'!\[[^\]]*\]\([^)]*\.(?:svg|png)\)', body):
                fail(file, 'ARCHITECTURE', 'formal system document requires its actual architecture diagram')
            if re.search(r'(?m)^\s*(?:sequenceDiagram|stateDiagram)', body):
                fail(file, 'FORMAL_DIAGRAM', 'formal system document should retain architecture, not copied planning diagrams')
        for match in re.finditer(r'\]\(([^)]+)\)', body):
            link(file, match[1], text[:offset + match.start()].count('\n') + 1)
        for graph in re.findall(r'```mermaid\s*\n(.*?)\n```', body, re.S):
            nodes = dict(re.findall(r'(\w+)\["([^"]+)"\]', graph))
            for node, ref in re.findall(r'click\s+(\w+)\s+"([^"]+)"', graph):
                link(file, ref)
                target = file.parent / ref.split('#')[0]
                if node not in nodes:
                    fail(file, 'GRAPH_NODE', 'linked node is not declared: ' + node)
                elif target.exists() and target.suffix == '.md':
                    label = nodes[node].split('<br/>')[0]
                    titles = re.findall(r'^#{1,6}\s+(.+)$', target.read_text(encoding='utf-8-sig'), re.M)
                    if label not in titles:
                        fail(file, 'GRAPH_NAME', 'node/document title mismatch: ' + label)
    if count == 0 and not issues:
        fail(path, 'EMPTY_SCOPE', 'no supported document/declaration files checked')
    output({'ok': not issues, 'checked_files': count, 'issues': issues,
            'skipped_template_count': len(skipped), 'limits': ['no business/ownership verification', 'no Mermaid rendering or language API semantic parsing', 'remote URLs not fetched'] + ([] if args.project_root else ['source existence not checked without --project-root'])})
    return 0 if not issues else 1


def config(args):
    path = Path(args.path)
    if not path.exists():
        output({'exists': False, 'next': 'plan configuration and field-processing list; do not create/export a table', 'path': str(path)})
        return 0
    if path.suffix.lower() == '.xls':
        import xlrd
        book = xlrd.open_workbook(str(path), on_demand=True)
        if args.sheet not in book.sheet_names():
            output({'exists': False, 'workbook_exists': True, 'sheets': book.sheet_names(), 'next': 'plan missing sheet, preserve existing workbook'})
            return 0
        sheet = book.sheet_by_name(args.sheet)
        rows = [sheet.row_values(i) for i in range(sheet.nrows)]
    elif path.suffix.lower() == '.xlsx':
        import openpyxl
        book = openpyxl.load_workbook(path, read_only=True, data_only=False)
        try:
            if args.sheet not in book.sheetnames:
                output({'exists': False, 'workbook_exists': True, 'sheets': book.sheetnames, 'next': 'plan missing sheet'})
                return 0
            rows = [[v if v is not None else '' for v in r] for r in book[args.sheet].iter_rows(values_only=True)]
        finally:
            book.close()
    else:
        raise ValueError('only .xls/.xlsx supported; no inferred parser for other formats')
    output(tagged_config(rows, args.sheet, args.limit, args.summary))
    return 0


def tagged_config(rows, sheet, limit=20, summary=False):
    tags = {}
    aliases = {'OUT_CLICENT': 'OUT_CLIENT', 'OUT_CLICNET': 'OUT_CLIENT'}
    for row in rows:
        if not row:
            continue
        key = str(row[0]).strip().upper()
        key = aliases.get(key, key)
        if key in {'OUT_CLIENT', 'OUT_SERVER', 'IS_CLIENT', 'IS_SERVER', 'KEY', 'TYPE', 'TIPS'} and key in tags:
            raise ValueError('duplicate table tag: ' + key)
        if key != 'VALUE':
            tags[key] = row
    if not {'KEY', 'TYPE'}.issubset(tags):
        raise ValueError('not a supported tagged workbook; use the project reader, do not guess header layout')
    fields, issues, seen = [], [], set()
    table_sides = {side: bool(str(tags.get('OUT_' + side.upper(), ['', ''])[1] if len(tags.get('OUT_' + side.upper(), [])) > 1 else '').strip())
                   and str(tags.get('OUT_' + side.upper(), ['', ''])[1]).strip().upper() != 'NO'
                   for side in ('client', 'server')}
    width = len(tags['KEY'])
    def cell(tag, index):
        row = tags.get(tag, [])
        return row[index] if index < len(row) else ''
    for i in range(1, width):
        name = str(cell('KEY', i)).strip()
        if not name:
            continue
        if name in seen:
            raise ValueError('duplicate field: ' + name)
        seen.add(name)
        tip = str(cell('TIPS', i)).strip()
        side_flags = {side: table_sides[side] and str(cell('IS_' + side.upper(), i)).upper() in {'YES', 'TRUE', '1'} for side in ('client', 'server')}
        if any(side_flags.values()) and not tip:
            issues.append({'field': name, 'rule': 'MISSING_TIPS'})
        fields.append({'name': name, 'source_type': cell('TYPE', i), 'tips': tip, 'column': i + 1, **side_flags})
    values = [{f['name']: row[f['column'] - 1] if f['column'] <= len(row) else '' for f in fields}
              for row in rows if row and str(row[0]).strip().upper() == 'VALUE']
    result = {'exists': True, 'sheet': sheet, 'table_sides': table_sides, 'next': 'read/reuse; plan only reported gaps',
            'field_count': len(fields), 'enabled_field_counts': {side: sum(f[side] for f in fields) for side in ('client', 'server')},
            'record_count': len(values),
            'issues': issues, 'limits': ['raw source values only; no guessed defaults/type conversions', 'does not execute exporter or runtime API', 'tagged format only']}
    if not summary:
        result.update(fields=fields, records=values[:limit], truncated=len(values) > limit)
    return result


def logs(args):
    patterns = [re.compile(r'(?:^|[\s,{])' + re.escape(k) + r'\s*[:=]\s*"?' + re.escape(v) + r'"?(?=[\s,}]|$)')
                for k, v in (item.split('=', 1) for item in args.match)]
    result, total = [], 0
    for file_name in args.files:
        path = Path(file_name)
        before = deque(maxlen=args.context)
        remaining, last_emitted = 0, 0
        with path.open(encoding=args.encoding, errors='strict') as stream:
            for number, line in enumerate(stream, 1):
                line = line.rstrip('\r\n')
                matched = all(p.search(line) for p in patterns) and (not args.contains or args.contains in line)
                if matched and total < args.limit:
                    total += 1
                    for n, text in before:
                        if n > last_emitted:
                            result.append({'file': str(path), 'line': n, 'text': text, 'match': False})
                            last_emitted = n
                    if number > last_emitted:
                        result.append({'file': str(path), 'line': number, 'text': line, 'match': True})
                        last_emitted = number
                    remaining = args.context
                elif remaining:
                    if number > last_emitted:
                        result.append({'file': str(path), 'line': number, 'text': line, 'match': False})
                        last_emitted = number
                    remaining -= 1
                before.append((number, line))
                if total >= args.limit and not remaining:
                    break
    output({'matches': total, 'limit_reached': total >= args.limit, 'evidence': result,
            'limits': ['does not infer session identity, timezone order or root cause', 'raw evidence; redact sensitive values before sharing']})
    return 0


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest='command', required=True)
    rules = sub.add_parser('rules', help='print compact structural contracts')
    rules.add_argument('--kind', choices=list(SCHEMAS), required=True)
    validate = sub.add_parser('check', help='read-only Markdown/API layout check; JSON diagnostics')
    validate.add_argument('path')
    validate.add_argument('--kind', choices=['auto'] + list(SCHEMAS), default='auto')
    validate.add_argument('--project-root')
    validate.add_argument('--allow-meta-field', action='append', default=[])
    validate.add_argument('--topic', action='store_true', help='formal UI/config topic: no separate architecture diagram requirement')
    inspect = sub.add_parser('config', help='read-only tagged .xls/.xlsx inspection or missing-config decision')
    inspect.add_argument('path')
    inspect.add_argument('--sheet', required=True)
    inspect.add_argument('--limit', type=int, default=20)
    inspect.add_argument('--summary', action='store_true', help='only counts/sides/issues; omit full field/value dump')
    read = sub.add_parser('logs', help='stream matching lines with context; no diagnosis')
    read.add_argument('files', nargs='+')
    read.add_argument('--match', action='append', default=[], help='exact key=value token, repeat for session and request')
    read.add_argument('--contains')
    read.add_argument('--context', type=int, default=4)
    read.add_argument('--limit', type=int, default=20)
    read.add_argument('--encoding', default='utf-8-sig')
    args = parser.parse_args()
    try:
        if hasattr(args, 'limit') and args.limit < 1:
            raise ValueError('--limit must be positive')
        if hasattr(args, 'context') and not 0 <= args.context <= 1000:
            raise ValueError('--context must be between 0 and 1000')
        if args.command == 'rules':
            output({'kind': args.kind, 'required_meta': SCHEMAS[args.kind], 'sides': sorted(SIDES), 'check_is_read_only': True})
            return 0
        return {'check': check, 'config': config, 'logs': logs}[args.command](args)
    except Exception as exc:
        output({'ok': False, 'error': str(exc), 'exception': type(exc).__name__})
        return 2


if __name__ == '__main__':
    raise SystemExit(main())
